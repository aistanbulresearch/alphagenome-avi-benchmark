"""Supplementary tables for the manuscript, built from the aggregate result tables in results/ and, for Supplementary
Table 3, from the label rules and effect measures in data/derived/part2/assay_variants.tsv.gz (no tool score is
read or written).

Usage (from the repository root): python manuscript/tools/make_supplementary_tables.py <output.xlsx>
Writes one workbook: a Contents sheet (number, title, legend, source, then definitions of sets, regions, tools and
abbreviations) and one sheet per table.
"""
from __future__ import annotations

import ast
import json
import re
import subprocess
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

R1, R2 = Path("results/part1"), Path("results/part2")
OUT = Path(sys.argv[1])
TITLE = ("Benchmarking the AlphaGenome Variant Impact score against SpliceAI on post-release ClinVar "
         "classifications and experimental splicing assays")
# Tables are built in topic order (positions 1-17 in tables()) and numbered by their first citation in the
# manuscript body: NUMBERING_ORDER lists the topic positions in that order.
NUMBERING_ORDER = [1, 2, 12, 3, 4, 6, 7, 8, 10, 9, 13, 14, 5, 11, 15, 16, 17]
SMALL = "fewer than 5 P/LP or B/LB"
SMALL2 = "fewer than 5 splice-altering or other"
DEF = "Sets, regions, tools and abbreviations are defined on the Contents sheet."

DATASETS = {
    "kitzman_pou1f1_ex2": "POU1F1 exon 2 (Gergics et al. 2021)",
    "kitzman_wt1_ex9": "WT1 exon 9 (Smith et al. 2023)",
    "braun2018_mst1r_ex11": "MST1R exon 11 (Braun et al. 2018)",
    "cortes_lopez2022_cd19": "CD19 exons 1–3 (Cortés-López et al. 2022)",
    "compass": "COMPASS (Koplik et al. 2025)",
    "opensplice": "OpenSplice (Quarantani et al. 2026)",
    "mfass": "MFASS (Chong et al. 2019), comparison dataset",
    "opensplice, developers' benchmark positions": "OpenSplice, exons in genes with the developers' saturation genome "
                                                   "editing benchmarks (reported separately)",
    "compass, random SNVs": "COMPASS without its ExAC subset",
    "compass, ExAC SNVs": "COMPASS ExAC subset (descriptive)",
}
ORDER = list(DATASETS)
SETS = {"main": "main set", "all frequencies": "all frequencies", "present and rare": "rare, present in gnomAD",
        "training chromosomes": "training chromosomes", "two stars": "at least two review stars",
        "stricter exclusion": "stricter exclusion of older unresolved records",
        "predictor mentioned": "submission mentioned a splicing predictor",
        "predictor not mentioned": "no submission mentioned a splicing predictor",
        "calibration sample": "calibration sample"}
REGIONS = {"intronic >=3 bp (main)": "main region (intronic, at least 3 bases from the exon)"}
TARGETS = {"supporting": 2.406, "moderate": 5.790, "strong": 33.53, "very_strong": 1124.0}

DEFINITIONS = [
    ("SETS (Part 1)", ""),
    ("main set", "ClinVar new labels (P/LP or B/LB in the October 2026 release and none in the 2025-06-15 release or the "
                 "July 2024 and July 2025 archives), after the exclusions in Supplementary Table 1: single-nucleotide "
                 "variants in intronic, untranslated or synonymous positions, with at least one review star, an HGVS "
                 "description on the MANE Select transcript, on the developers' validation or untouched chromosomes, "
                 "and rare in gnomAD v4.1 (FAF95 below 0.001, or absent)."),
    ("pooled test set", "The main set's non-canonical intronic and untranslated variants; the calibration thresholds are "
                        "tested on it (461 P/LP and 28,858 B/LB for AVI)."),
    ("all frequencies", "The main set without the rarity filter."),
    ("rare, present in gnomAD", "Main-set variants that have a gnomAD v4.1 record."),
    ("training chromosomes", "New labels passing the main-set filters, including rarity, but on the developers' "
                             "training chromosomes (1, 4, 7, 8, 10, 13, 15)."),
    ("at least two review stars", "Main-set variants with at least two ClinVar review stars."),
    ("stricter exclusion of older unresolved records", "The main set under the earlier rule that also excluded variants "
                                                       "inside longer older ClinVar records of unresolved identity."),
    ("submission mentioned / no submission mentioned a splicing predictor",
     "Main-set variants split by whether any ClinVar submission names a splicing predictor (fixed list in the "
     "repository; presence only)."),
    ("calibration sample", "Labelled variants outside the new-label set: non-canonical intronic and untranslated "
                           "variants with at least one star, every P/LP variant and 40,000 B/LB variants drawn at "
                           "random, stratified by region and chromosome group, then restricted to rare variants (for "
                           "AVI, 3,248 P/LP and 23,327 B/LB remain); thresholds are derived on it."),
    ("REGIONS (MANE Select HGVS offsets)", ""),
    ("canonical", "±1 and ±2."),
    ("donor region", "+3 to +6."),
    ("acceptor region", "−3 to −20."),
    ("near intronic", "+7 to +50 and −21 to −50."),
    ("deep intronic", "More than 50 bases from the exon."),
    ("main region", "Intronic, at least 3 bases from the exon: donor, acceptor, near intronic and deep intronic pooled."),
    ("UTR, 5' UTR, 3' UTR; synonymous", "Untranslated variants; synonymous coding variants."),
    ("TOOLS", ""),
    ("AVI", "Raw score of the AVI scorer (AVI_SCORE), AlphaGenome Atlas interface, Python client 0.9.0."),
    ("SpliceAI", "Precomputed raw (unmasked) scores for Ensembl MANE v1.4, 50-base window; maximum of the four delta "
                 "scores and over gene entries."),
    ("SpliceAI masked", "The same precomputed files in their masked form (changes that strengthen annotated or weaken "
                        "unannotated splice sites set to zero)."),
    ("SpliceAI 4999", "SpliceAI 1.3.1 run locally with a 4,999-base window, deep intronic variants only."),
    ("SpliceAI2", "Precomputed summary score, version 2.0; a comparator added after the main results under a rule "
                  "fixed before its scores were retrieved; research use only (no per-variant score is given)."),
    ("Pangolin", "Pangolin 1.0.2 run locally (GENCODE v38, masked, 50-base window). Part 1: every non-canonical "
                 "intronic and untranslated P/LP of the main set (461) and a region-stratified sample of 4,000 B/LB; "
                 "every metric uses inverse-sampling weights. Part 2: the four minigene datasets only."),
    ("CADD", "CADD v1.7 PHRED score."),
    ("phyloP 100-way; phastCons 470-way; phyloP 241-way", "Conservation scores (UCSC hg38); phastCons 470-way and "
                                                         "Cactus 241-way phyloP are AVI's own conservation inputs."),
    ("Distance", "Distance to the nearest exon boundary from the HGVS offset; smaller distances ranked as more likely "
                 "pathogenic (planned baseline, reported on its own)."),
    ("ABBREVIATIONS", ""),
    ("P/LP; B/LB", "Pathogenic or likely pathogenic; benign or likely benign."),
    ("SNV; UTR; MANE; HGVS", "Single-nucleotide variant; untranslated region; Matched Annotation from NCBI and EMBL-EBI; "
                             "Human Genome Variation Society nomenclature."),
    ("FAF95", "gnomAD joint filtering allele frequency (fafmax_faf95_max_joint)."),
    ("AUROC; AUPRC", "Area under the receiver operating characteristic curve; area under the precision-recall curve."),
    ("LR", "Likelihood ratio. ClinGen targets (prior 0.0441): supporting 2.406, moderate 5.790, strong 33.53, very "
           "strong 1,124; benign targets are the reciprocals."),
    ("DL; HK", "DerSimonian–Laird random-effects pooling; Hartung–Knapp interval."),
    ("I²; tau²; k", "Heterogeneity share; between-dataset variance; number of datasets pooled."),
    ("SDV; PSI", "Splice-disrupting variant; percent spliced in (exon inclusion)."),
    ("BH; FDR", "Benjamini–Hochberg; false discovery rate."),
    ("KTS", "WT1 isoforms with (+KTS) or without (−KTS) the lysine–threonine–serine insertion."),
    ("SGE", "Saturation genome editing."),
    ("NOTATION", ""),
    ("intervals", "95% intervals. 'no upper bound' and 'no lower bound' mark an interval (of scores or of a likelihood "
                  "ratio) open above or below; 'not reached' marks a threshold that the data did not reach. Metrics "
                  "are blank where a class (P/LP or B/LB) is empty."),
    (SMALL, "Flag for estimates resting on fewer than 5 P/LP or 5 B/LB (Part 2: splice-altering or other); such "
            "estimates are shown but carry little information."),
    ("rounding", "Metrics and differences to four decimals; likelihood ratios and thresholds to four significant "
                 "figures."),
]


def ci(value) -> tuple:
    """(low, high) from '(a, b)', '[a, b]' or a list."""
    if isinstance(value, (list, tuple)):
        lo, hi = value
    elif isinstance(value, str) and value.strip():
        lo, hi = ast.literal_eval(value.replace("nan", "None").replace("-inf", "-1e999").replace("inf", "1e999"))
    else:
        return None, None
    return lo, hi


def bound(x):
    if x == float("inf") or (isinstance(x, float) and x > 1e300):
        return "no upper bound"
    if x == float("-inf") or (isinstance(x, float) and x < -1e300):
        return "no lower bound"
    return None


def r(x, d=4):
    if x is None:
        return None
    try:
        x = float(x)
    except (TypeError, ValueError):
        return x
    if x != x:
        return None
    return bound(x) or round(x, d)


def sig(x, n=4):
    if x is None:
        return None
    x = float(x)
    if x != x:
        return None
    return bound(x) or float(f"{x:.{n}g}")


def small(a, b) -> str:
    try:
        return "yes" if int(a) < 5 or int(b) < 5 else ""
    except (TypeError, ValueError):
        return ""


def symbols(t: str) -> str:
    return str(t).replace(">=", "≥").replace("<=", "≤")


def rule_text(t: str) -> str:
    """Label rules and effect measures: ≥/≤, and a true minus sign before numbers and in 'KTS-'."""
    t = symbols(t)
    t = re.sub(r"([≤≥<>]\s?)-(\d)", "\\1\u2212\\2", t)
    return t.replace("KTS-", "KTS\u2212")


def ds(name: str) -> str:
    return DATASETS.get(name, name)


def supported(level: str, lo, hi) -> tuple:
    target = TARGETS["supporting" if level.startswith("published") else level.split()[0]]
    if "pathogenic" in level:
        return target, ("yes" if isinstance(lo, (int, float)) and lo > target else "no")
    return round(1 / target, 4), ("yes" if isinstance(hi, (int, float)) and hi < 1 / target else "no")


def tables() -> list:
    out = []
    # 1 flow
    f = pd.read_csv(R1 / "flow.tsv", sep="\t").fillna("")
    f["step"] = f["step"].map(symbols)
    out.append(("Part 1: flow from the ClinVar October 2026 release to the main set",
                "Variants (P/LP, B/LB, total) at each step. Exclusions are applied in the order listed. Variants on "
                "training chromosomes are counted before the rarity filter. " + DEF, "results/part1/flow.tsv", f))
    # 2 coverage
    c = pd.read_csv(R1 / "coverage.tsv", sep="\t")
    c["set"] = c["set"].map(lambda s: SETS.get(s, s))
    c["fraction scored"] = (c["scored"] / c["variants"]).round(4)
    notes = {"Pangolin": "run on a sample only (461 P/LP and 4,000 B/LB of the main set's non-canonical intronic and "
                         "untranslated variants); the fraction is of the whole set",
             "SpliceAI 4999": "run on deep intronic variants only; the fraction is of the whole set",
             "Distance": "intronic variants with an HGVS offset only"}
    c["note"] = c["tool"].map(lambda t: notes.get(t, ""))
    out.append(("Part 1: coverage per tool and set",
                "Variants in each set and the number each tool scored. Tools were run on every variant unless the note "
                "says otherwise. " + DEF, "results/part1/coverage.tsv", c))
    # 3 discrimination
    d = pd.read_csv(R1 / "discrimination.tsv", sep="\t")
    rows, dropped = [], 0
    for _, x in d.iterrows():
        if (pd.isna(x["n_pos"]) or int(x["n_pos"]) == 0) and (pd.isna(x["n_neg"]) or int(x["n_neg"]) == 0):
            dropped += 1
            continue
        a_lo, a_hi = ci(x["auroc_ci"])
        p_lo, p_hi = ci(x["auprc_ci"])
        rows.append({"set": SETS.get(x["set"], x["set"]), "region": REGIONS.get(x["region"], x["region"]),
                     "tool": x["tool"], "P/LP": x["n_pos"], "B/LB": x["n_neg"], "genes with P/LP": x["genes_pos"],
                     "AUROC": r(x["auroc"]), "AUROC 95% low": r(a_lo), "AUROC 95% high": r(a_hi),
                     "AUPRC": r(x["auprc"]), "AUPRC 95% low": r(p_lo), "AUPRC 95% high": r(p_hi),
                     "AUPRC baseline (P/LP fraction)": r(x["pos_fraction"]), SMALL: small(x["n_pos"], x["n_neg"])})
    print(f"  discrimination: {dropped} rows without scored variants omitted")
    out.append(("Part 1: discrimination by set, region and tool",
                "AUROC and AUPRC with 95% intervals from 2,000 gene-bootstrap replicates (seed 20261007). Pangolin "
                "metrics use inverse-sampling weights; its B/LB count is the scored sample. Rows for a tool that was "
                f"not run on a set or region (for example Pangolin outside the main set's sample, SpliceAI 4999 outside "
                f"deep intronic variants) are omitted ({dropped} rows); rows where the tool was run but one class is "
                f"empty are kept, with blank metrics. " + DEF, "results/part1/discrimination.tsv",
                pd.DataFrame(rows)))
    # 4 paired AUROC differences
    p = pd.read_csv(R1 / "paired_differences.tsv", sep="\t")
    rows = []
    for _, x in p.iterrows():
        lo, hi = ci(x["ci"])
        rows.append({"set": SETS.get(x["set"], x["set"]), "region": REGIONS.get(x["region"], x["region"]),
                     "comparison": x["pair"], "P/LP": x["n_pos"], "B/LB": x["n_neg"],
                     "AUROC difference": r(x["difference"]), "95% low": r(lo), "95% high": r(hi),
                     SMALL: small(x["n_pos"], x["n_neg"])})
    out.append(("Part 1: paired AUROC differences, AVI minus each tool, in every set and region",
                "Differences on the variants both tools scored, with gene-bootstrap 95% intervals (not adjusted for "
                "multiple comparisons). The main comparison is AVI minus SpliceAI in the main region of the main set; "
                "the other sets are the supporting analyses. Comparisons with Pangolin use inverse-sampling weights. "
                "AVI minus masked SpliceAI was not computed in Part 1 (masked scores were a secondary form, used in "
                "calibration); it is reported for Part 2. " + DEF, "results/part1/paired_differences.tsv",
                pd.DataFrame(rows)))
    # 5 AUPRC and sensitivity
    s = pd.read_csv(R1 / "supplementary_discrimination.tsv", sep="\t")
    rows = []
    for _, x in s.iterrows():
        lo, hi = ci(x["ci"])
        rows.append({"set": SETS.get(x["set"], x["set"]), "region": REGIONS.get(x["region"], x["region"]),
                     "comparison": x["pair"], "metric": x["metric"], "P/LP": x["n_pos"], "B/LB": x["n_neg"],
                     "AVI": r(x["value_a"]), "other tool": r(x["value_b"]), "difference": r(x["difference"]),
                     "95% low": r(lo), "95% high": r(hi), SMALL: small(x["n_pos"], x["n_neg"])})
    out.append(("Part 1: paired AUPRC differences and sensitivity at 95% and 99% specificity (additional analysis)",
                "Added after the main results. Gene-bootstrap 95% intervals; Pangolin values use inverse-sampling "
                "weights. " + DEF, "results/part1/supplementary_discrimination.tsv", pd.DataFrame(rows)))
    cal = json.load(open(R1 / "calibration.json", encoding="utf-8"))
    # 6 thresholds
    rows = []
    for tool, t in cal.items():
        for level, (path, ben) in t["derivation"]["thresholds"].items():
            rows.append({"tool": tool, "derivation P/LP": t["derivation"]["n_pos"],
                         "derivation B/LB": t["derivation"]["n_neg"], "level": level.replace("_", " "),
                         "target LR (pathogenic)": TARGETS[level], "target LR (benign)": round(1 / TARGETS[level], 4),
                         "pathogenic threshold (score ≥)": sig(path) if path is not None else "not reached",
                         "benign threshold (score ≤)": sig(ben) if ben is not None else "not reached"})
    out.append(("Calibration: thresholds derived on the calibration sample (prior 0.0441)",
                "Thresholds from the procedure adapted from Pejaver et al. 2022 (Methods). For CADD these are the "
                "intronic thresholds referred to in the text. Where a moderate threshold equals the supporting one "
                "(phastCons 470-way, whose score saturates at 1), the supporting interval is empty. " + DEF,
                "results/part1/calibration.json", pd.DataFrame(rows)))
    rows, rows_cum, rows_reg, rows_prior = [], [], [], []
    for tool, t in cal.items():
        e = t["evaluated_on_main"]
        # a level whose score interval has equal ends (phastCons 470-way saturates at 1) holds no variant at all
        empty = {lvl for lvl, v in e["interval"].items()
                 if v.get("interval") and v["interval"][0] is not None and v["interval"][0] == v["interval"][1]}
        for level, v in e["interval"].items():
            lo, hi = ci(v["ci"])
            target, ok = supported(level, lo, hi)
            iv_lo, iv_hi = v.get("interval", [None, None])
            rows.append({"tool": tool, "test P/LP": t["test_n_pos"], "test B/LB": t["test_n_neg"], "level": level,
                         "score interval from": sig(iv_lo),
                         "score interval to": sig(iv_hi) if iv_hi is not None else "no upper bound",
                         "likelihood ratio": sig(v["lr"]), "95% low": sig(lo), "95% high": sig(hi),
                         "target LR": target, "supported": ok,
                         "interval method": "empty score interval" if level in empty else v.get("ci_method"),
                         "P/LP in interval": v.get("n_pos"), "B/LB in interval": v.get("n_neg"),
                         "share of P/LP": r(v.get("fraction_pos")), "share of B/LB": r(v.get("fraction_neg"))})
        for level, v in e["cumulative"].items():
            lo, hi = ci(v["ci"])
            target, ok = supported(level, lo, hi)
            rows_cum.append({"tool": tool, "level": symbols(level), "threshold": sig(v.get("threshold")),
                             "likelihood ratio of all variants beyond the threshold": sig(v["lr"]),
                             "95% low": sig(lo), "95% high": sig(hi), "target LR": target, "supported": ok,
                             "interval method": v.get("ci_method"),
                             "share of P/LP meeting": r(v.get("fraction_pos_meeting")),
                             "share of B/LB meeting": r(v.get("fraction_neg_meeting"))})
        for region, rv in e["interval_by_region"].items():
            for level, v in rv["levels"].items():
                lo, hi = ci(v["ci"])
                target, ok = supported(level, lo, hi)
                rows_reg.append({"tool": tool, "region": region, "region P/LP": rv["n_pos"], "region B/LB": rv["n_neg"],
                                 "level": level, "likelihood ratio": sig(v.get("lr")), "95% low": sig(lo),
                                 "95% high": sig(hi), "target LR": target, "supported": ok,
                                 "interval method": "empty score interval" if level in empty else
                                 ("no variant in the interval" if v.get("n_pos") == 0 and v.get("n_neg") == 0
                                  else v.get("ci_method")),
                                 "P/LP in interval": v.get("n_pos"), "B/LB in interval": v.get("n_neg")})
        for prior, pv in e["prior_sensitivity"].items():
            for level, (path, ben) in pv["thresholds"].items():
                rows_prior.append({"tool": tool, "prior": float(prior), "level": level.replace("_", " "),
                                   "target LR (pathogenic)": sig(pv["targets"][level]),
                                   "target LR (benign)": sig(1 / pv["targets"][level]),
                                   "pathogenic threshold (score ≥)": sig(path) if path is not None else "not reached",
                                   "benign threshold (score ≤)": sig(ben) if ben is not None else "not reached",
                                   "share of P/LP with at least supporting evidence": r(pv.get("fraction_pos_at_least_supporting")),
                                   "share of B/LB with supporting benign evidence": r(pv.get("fraction_neg_benign_supporting"))})
    exact = ("Intervals from a gene bootstrap, or, where the interval holds no P/LP or no B/LB, exact Clopper–Pearson "
             "bounds of the two proportions, each at 97.5% so that the ratio interval covers at least 95% "
             "(Bonferroni).")
    out.append(("Calibration on the pooled test set: likelihood ratio within each evidence level's score interval",
                "Pooled test set: the main set's non-canonical intronic and untranslated variants. A level is "
                "supported when its 95% interval lies above the target (below the benign target for benign levels). "
                + exact + " " + DEF, "results/part1/calibration.json", pd.DataFrame(rows)))
    out.append(("Calibration on the pooled test set: likelihood ratio of all variants beyond each threshold",
                "The cumulative form specified in the initial plan, on the pooled test set; 'supported' as in the "
                "interval form. Verdicts agree with the interval form for AVI, SpliceAI, masked SpliceAI, SpliceAI2, "
                "CADD and phyloP 100-way; at the supporting pathogenic level of phyloP 241-way and phastCons 470-way, "
                "whose supporting intervals are narrow or empty, the cumulative form is supported and the interval "
                "form is not. SpliceAI rows marked 'published' use the published ClinGen thresholds. " + exact + " "
                + DEF, "results/part1/calibration.json", pd.DataFrame(rows_cum)))
    out.append(("Calibration by region on the main set",
                "Interval likelihood ratios within each region of the pooled test set; Table 2 of the main text is a "
                "subset. Untranslated regions hold 12 P/LP. " + exact + " " + DEF,
                "results/part1/calibration.json", pd.DataFrame(rows_reg)))
    out.append(("Calibration under other priors (0.01 and 0.10)",
                "Targets from the formula of Pejaver et al. 2022 at each prior; thresholds rederived on the "
                "calibration sample; the shares are computed on the pooled test set. " + DEF,
                "results/part1/calibration.json", pd.DataFrame(rows_prior)))
    # 11 published thresholds
    pt = pd.read_csv(R1 / "published_thresholds.tsv", sep="\t")
    rows = []
    for _, x in pt.iterrows():
        lo, hi = ci(x["ci"])
        level = ("pathogenic " if str(x["level"]).startswith("PP3") else "benign ") + "supporting"
        rows.append({"tool and source": x["tool"], "level": x["level"],
                     "score from": bound(float(x["lo"])) or float(x["lo"]), "score to": bound(float(x["hi"])) or float(x["hi"]),
                     "side": {"ge": "score ≥ threshold", "le": "score ≤ threshold"}.get(x["side"], x["side"]),
                     "likelihood ratio": sig(x["lr"]), "95% low": sig(lo), "95% high": sig(hi),
                     "P/LP": x["n_pos"], "B/LB": x["n_neg"], "share of P/LP": r(x["fraction_pos"]),
                     "share of B/LB": r(x["fraction_neg"])})
    out.append(("Published thresholds on the pooled test set (non-canonical intronic and untranslated variants of the "
                "main set; additional analysis)",
                "SpliceAI ClinGen thresholds (Walker et al. 2023) applied to the 50-base precomputed scores, and the "
                "CADD score intervals recommended for missense variants (Pejaver et al. 2022). " + DEF,
                "results/part1/published_thresholds.tsv", pd.DataFrame(rows)))
    # 12 Part 2 datasets
    pj = json.load(open(R2 / "pooled.json", encoding="utf-8"))
    av = pd.read_csv("data/derived/part2/assay_variants.tsv.gz", sep="\t", dtype=str, keep_default_na=False,
                     usecols=["dataset", "threshold_rule", "use", "splice_altering", "effect_measure", "subset",
                              "variant_type", "chrom", "pos", "ref", "alt"])

    def one_measure(values) -> str:
        """Collapse measures that differ only in their '(cells: ...)' suffix into one phrase."""
        cells, bases = [], set()
        for m in values:
            hit = re.match(r"^(.*?)\s*\(cells: ([^)]*)\)$", m)
            if not hit:
                bases.add(m)
                continue
            bases.add(hit.group(1))
            cells += [c.strip() for c in hit.group(2).split(",")]
        if len(bases) == 1 and cells:
            base = next(iter(bases)).replace("mean over cell lines of pooled-replicate values",
                                             "mean of pooled-replicate values over the cell lines in which each "
                                             "variant was measured")
            order = ["HEK", "HeLa", "K562", "MCF7", "HMC3"]
            return f"{base} ({', '.join(sorted(set(cells), key=lambda c: order.index(c) if c in order else 99))})"
        return "; ".join(sorted(bases))

    def duplicates(snv: pd.DataFrame) -> tuple[int, int]:
        """(SNV rows set aside as duplicates, alleles whose records carry different calls)."""
        key = snv["chrom"] + ":" + snv["pos"] + ":" + snv["ref"] + ":" + snv["alt"]
        is_dup = snv["subset"].str.contains("duplicate")
        dup_keys = set(key[is_dup])
        calls = snv[key.isin(dup_keys)].assign(_k=key[key.isin(dup_keys)]).groupby("_k")["splice_altering"].nunique()
        return int(is_dup.sum()), int((calls > 1).sum())

    def set_aside(name: str, sub: pd.DataFrame) -> str:
        """SNVs not used by design or data rules before the counts below."""
        snv = sub[sub["variant_type"] == "SNV"]
        n_dup, n_diff = duplicates(snv)
        dup_text = f"{n_dup:,} duplicates ({n_diff:,} alleles with different calls in their records)"
        if name == "compass":
            nu = snv[snv["use"] == "not_used"]
            n_clin, n_geu = int((nu["subset"] == "clinvar").sum()), int((nu["subset"] == "geuvadis").sum())
            n_psi = int(nu["subset"].str.contains("reference_psi_0_or_1").sum())
            return (f"ClinVar subset ({n_clin:,}) and Geuvadis subset ({n_geu:,}); {n_psi + n_dup:,} other SNVs: "
                    f"{n_psi:,} with a reference PSI of 0 or 1 in every cell line in which it was measured and "
                    f"{dup_text}")
        if name.startswith("opensplice, "):
            return "see OpenSplice"
        return (f"{n_dup:,} duplicate SNVs ({n_diff:,} alleles with different calls in their records)"
                if n_dup else "none")

    # the exclusions of Methods, with the masks of src/part2_analysis.py (labels only, no scores)
    sys.path.insert(0, "src")
    import part2_analysis as p2a
    xd = p2a._read(p2a.TABLE)
    xd = xd[xd["use"].isin(["main", "comparison"]) & (xd["variant_type"] == "SNV")
            & xd["splice_altering"].isin(["yes", "no"])].copy()
    xd["allele"] = xd["chrom"] + ":" + xd["pos"] + ":" + xd["ref"] + ":" + xd["alt"]
    xd["in_part1"] = xd["allele"].isin(p2a.part1_new_labels())
    xd["development"] = xd["allele"].isin(p2a.development_hits(xd.drop_duplicates("allele")))
    xd["median"] = xd["subset"] == "median_estimate"

    def exclusions(name: str) -> dict:
        if name.startswith("opensplice, "):
            g = xd[(xd["dataset"] == "opensplice") & xd["development"]]
            return {"binary": len(g), "part1": "not applied", "dev": "this row is that set", "median": "none"}
        g = xd[xd["dataset"] == name]
        dev = g["development"] & ~g["in_part1"]
        med = g["median"] & ~g["in_part1"] & ~g["development"]
        dev_n = int(dev.sum())
        return {"binary": len(g), "part1": int(g["in_part1"].sum()),
                "dev": f"{dev_n:,} (reported separately, last row)" if name == "opensplice" else dev_n,
                "median": int(med.sum()) if name == "braun2018_mst1r_ex11" else "none"}

    rows = []
    for name in ["kitzman_pou1f1_ex2", "kitzman_wt1_ex9", "braun2018_mst1r_ex11", "cortes_lopez2022_cd19", "compass",
                 "opensplice", "mfass", "opensplice, developers' benchmark positions"]:
        base = "opensplice" if name.startswith("opensplice") else name
        sub = av[av["dataset"] == base]
        used = sub[sub["use"].isin(["main", "comparison"])]
        rules = sorted(set(used["threshold_rule"]))
        kept = [x for x in rules if not x.startswith("no test possible")]
        dropped_rules = [x for x in rules if x.startswith("no test possible")]
        measures = one_measure(sorted({m for m in used["effect_measure"] if m}))
        n_int = int((sub["splice_altering"] == "intermediate").sum())
        cnt = pj["counts"].get(name, {})
        ex = exclusions(name)
        rows.append({"dataset": ds(name), "role": "comparison" if name == "mfass" else
                     ("reported separately" if "," in name else "main"),
                     "set aside before the counts (design or data rules)": set_aside(name, sub),
                     "binary-labelled SNVs": ex["binary"],
                     "removed: Part 1 new labels": ex["part1"],
                     "removed: developers' benchmark data": ex["dev"],
                     "removed: median-based estimates without a test": ex["median"],
                     "analysed variants": cnt.get("variants"), "splice-altering": cnt.get("splice_altering"),
                     "label rule (authors' calls or the authors' threshold)": rule_text("; ".join(kept)),
                     "effect measure (used for the Spearman correlation)": rule_text(measures),
                     "intermediate variants left out of binary metrics": n_int if n_int else "none defined"})
    out.append(("Part 2 datasets: label definitions and analysed counts",
                "Binary-labelled SNVs are the usable single-nucleotide variants with a splice-altering or other call "
                "(intermediate variants not included). The analysed count is that number minus the exclusions of "
                "Methods: Part 1 new labels, the developers' benchmark data and, for MST1R, median-based estimates "
                "without a test (the OpenSplice variants in the developers' benchmark genes are reported separately, "
                "last row). COMPASS enters with its random variants and its ExAC subset; the SNVs it sets aside by "
                "design or data rules are not in its binary-labelled count. Intermediate counts are before the "
                "exclusions. " + DEF, "results/part2/pooled.json; data/derived/part2/assay_variants.tsv.gz (label "
                "rules, effect measures and the exclusion masks of src/part2_analysis.py; no scores)", pd.DataFrame(rows)))
    # 13 Part 2 discrimination
    d2 = pd.read_csv(R2 / "discrimination.tsv", sep="\t")
    d2["_o"] = d2["dataset"].map(lambda n: ORDER.index(n) if n in ORDER else 99)
    rows, dropped2 = [], 0
    for _, x in d2.sort_values(["_o"], kind="stable").iterrows():
        if (pd.isna(x["n_pos"]) or int(x["n_pos"]) == 0) and (pd.isna(x["n_neg"]) or int(x["n_neg"]) == 0):
            dropped2 += 1
            continue
        a_lo, a_hi = ci(x["auroc_ci"])
        p_lo, p_hi = ci(x["auprc_ci"])
        rows.append({"dataset": ds(x["dataset"]), "tool": x["tool"], "splice-altering": x["n_pos"], "other": x["n_neg"],
                     "AUROC": r(x["auroc"]), "AUROC 95% low": r(a_lo), "AUROC 95% high": r(a_hi),
                     "AUPRC": r(x["auprc"]), "AUPRC 95% low": r(p_lo), "AUPRC 95% high": r(p_hi),
                     "AUPRC baseline": r(x["pos_fraction"]), "Spearman with |effect|": r(x["spearman_abs_effect"]),
                     SMALL2: small(x["n_pos"], x["n_neg"])})
    out.append(("Part 2: discrimination per dataset and tool",
                "95% intervals from 2,000 bootstrap replicates resampling variants (single-gene datasets) or genes "
                "(COMPASS, OpenSplice, MFASS). Spearman with |effect|: rank correlation of each score with the "
                "absolute effect size (the dataset's effect measure, listed with the Part 2 datasets), on the variants "
                f"the tool scored. Pangolin was run on the four minigene datasets only; its rows for the other datasets "
                f"are omitted ({dropped2} rows). " + DEF,
                "results/part2/discrimination.tsv", pd.DataFrame(rows)))
    # 14 Part 2 paired + pooled
    p2 = pd.read_csv(R2 / "paired_differences.tsv", sep="\t")
    p2["_o"] = p2["dataset"].map(lambda n: ORDER.index(n) if n in ORDER else 99)
    rows = []
    for _, x in p2.sort_values(["pair", "_o"], kind="stable").iterrows():
        lo, hi = ci(x["ci"])
        rows.append({"comparison": x["pair"], "dataset or pooled estimate": ds(x["dataset"]),
                     "splice-altering": x["n_pos"], "other": x["n_neg"], "AUROC difference": r(x["difference"]),
                     "95% low": r(lo), "95% high": r(hi), "Hartung–Knapp low": None, "Hartung–Knapp high": None,
                     "I²": None, "tau²": None, "k": None, SMALL2: small(x["n_pos"], x["n_neg"])})
    for pair, v in pj["pooled"].items():
        hk = v.get("ci_hartung_knapp") or [None, None]
        rows.append({"comparison": pair, "dataset or pooled estimate": "pooled (DerSimonian–Laird): " +
                     ", ".join(ds(n).split(" (")[0] for n in v["datasets"]),
                     "AUROC difference": r(v["estimate"]), "95% low": r(v["ci"][0]), "95% high": r(v["ci"][1]),
                     "Hartung–Knapp low": r(hk[0]), "Hartung–Knapp high": r(hk[1]), "I²": r(v["i2"], 3),
                     "tau²": r(v["tau2"], 5), "k": v["k"]})
        for left, lv in v.get("leave_one_out", {}).items():
            lhk = lv.get("ci_hartung_knapp") or [None, None]
            rows.append({"comparison": pair, "dataset or pooled estimate": f"pooled, leaving out {ds(left).split(' (')[0]}",
                         "AUROC difference": r(lv["estimate"]), "95% low": r(lv["ci"][0]), "95% high": r(lv["ci"][1]),
                         "Hartung–Knapp low": r(lhk[0]), "Hartung–Knapp high": r(lhk[1])})
    out.append(("Part 2: paired AUROC differences per dataset, pooled estimates and leave-one-dataset-out",
                "Random-effects pooling (DerSimonian–Laird) over the main datasets, with the bootstrap standard "
                "deviation as the standard error; Hartung–Knapp intervals and leave-one-out estimates were added "
                "after the main results. SpliceAI2 is pooled without OpenSplice, on which its developers evaluated it. "
                + DEF, "results/part2/paired_differences.tsv; results/part2/pooled.json", pd.DataFrame(rows)))
    # 15 by position: counts and per-dataset differences
    vp = pd.read_csv(R2 / "variant_positions.tsv", sep="\t")
    bd = pd.read_csv(R2 / "by_position_differences.tsv", sep="\t")
    rows = []
    for _, x in vp.iterrows():
        rows.append({"dataset": ds(x["dataset"]), "position class": x["class"], "variants": x["variants"],
                     "splice-altering": x["splice_altering"], "comparison": "", "AUROC difference": None,
                     "95% low": None, "95% high": None, SMALL2: ""})
    for _, x in bd.iterrows():
        lo, hi = ci(x["ci"])
        rows.append({"dataset": ds(x["dataset"]), "position class": x["class"], "variants": x["n_pos"] + x["n_neg"],
                     "splice-altering": x["n_pos"], "comparison": x["pair"], "AUROC difference": r(x["difference"]),
                     "95% low": r(lo), "95% high": r(hi), SMALL2: small(x["n_pos"], x["n_neg"])})
    out.append(("Part 2 by position relative to the exon: variant counts and per-dataset differences (additional analysis)",
                "Rows without a comparison give the variant counts per class; the others give AVI minus each tool "
                "within a class. Distances follow each assay's construct description; in CD19 and MST1R intronic "
                "distance is to the nearest exon of the construct. A dataset's class enters the pool when it holds "
                "at least 10 splice-altering and 10 other variants. " + DEF,
                "results/part2/variant_positions.tsv; results/part2/by_position_differences.tsv", pd.DataFrame(rows)))
    bp = json.load(open(R2 / "by_position_pooled.json", encoding="utf-8"))
    rows = []
    for cls, pairs in bp.items():
        for pair, v in pairs.items():
            hk = v.get("ci_hartung_knapp") or [None, None]
            rows.append({"position class": cls, "comparison": pair, "estimate": "pooled", "AUROC difference": r(v["estimate"]),
                         "95% low": r(v["ci"][0]), "95% high": r(v["ci"][1]), "Hartung–Knapp low": r(hk[0]),
                         "Hartung–Knapp high": r(hk[1]), "I²": r(v.get("i2"), 3), "k": v.get("k"),
                         "datasets": ", ".join(ds(n).split(" (")[0] for n in v.get("datasets", []))})
            for left, lv in v.get("leave_one_out", {}).items():
                lhk = lv.get("ci_hartung_knapp") or [None, None]
                rows.append({"position class": cls, "comparison": pair,
                             "estimate": f"leaving out {ds(left).split(' (')[0]}", "AUROC difference": r(lv["estimate"]),
                             "95% low": r(lv["ci"][0]), "95% high": r(lv["ci"][1]),
                             "Hartung–Knapp low": r(lhk[0]), "Hartung–Knapp high": r(lhk[1])})
    out.append(("Part 2 by position relative to the exon: pooled estimates (additional analysis)",
                "DerSimonian–Laird estimates with Hartung–Knapp intervals and leave-one-out estimates per class. "
                "Each class pools a different set of datasets; classes were not tested against each other. " + DEF,
                "results/part2/by_position_pooled.json", pd.DataFrame(rows)))
    # 17 COMPASS without ExAC
    cd = pd.read_csv(R2 / "compass_subsets_discrimination.tsv", sep="\t")
    cp = pd.read_csv(R2 / "compass_subsets_paired_differences.tsv", sep="\t")
    cj = json.load(open(R2 / "compass_subsets_pooled.json", encoding="utf-8"))
    cc = json.load(open(R2 / "compass_clinvar_check.json", encoding="utf-8"))
    rows = []
    for _, x in cd.iterrows():
        lo, hi = ci(x["auroc_ci"])
        rows.append({"part": ds(x["dataset"]), "quantity": f"AUROC, {x['tool']}", "splice-altering": x["n_pos"],
                     "other": x["n_neg"], "value": r(x["auroc"]), "95% low": r(lo), "95% high": r(hi)})
    for _, x in cp.iterrows():
        lo, hi = ci(x["ci"])
        rows.append({"part": ds(x["dataset"]), "quantity": f"AUROC difference, {x['pair']}", "splice-altering": x["n_pos"],
                     "other": x["n_neg"], "value": r(x["difference"]), "95% low": r(lo), "95% high": r(hi)})
    for pair, v in cj["pooled_with_random_compass"].items():
        hk = v.get("ci_hartung_knapp") or [None, None]
        rows.append({"part": "pooled over the main datasets with COMPASS without its ExAC subset",
                     "quantity": f"AUROC difference, {pair}", "value": r(v["estimate"]), "95% low": r(v["ci"][0]),
                     "95% high": r(v["ci"][1]), "Hartung–Knapp low": r(hk[0]), "Hartung–Knapp high": r(hk[1]),
                     "k": v["k"]})
    rel = {"variant_summary_2024-07.txt.gz": "July 2024", "variant_summary_2025-07.txt.gz": "July 2025",
           "variant_summary_2026-10.txt.gz": "October 2026"}
    for key, label in (("present_in", "all variants"), ("analysed_present_in", "analysed variants")):
        for f_, cnt in cc[key].items():
            rows.append({"part": "ClinVar check (labels only, no scores)",
                         "quantity": f"{label} present in the ClinVar summary of {rel[f_]}: random-labelled "
                                     f"{cnt.get('random_snv', 0):,}; ExAC subset {cnt.get('exac', 0):,}"})
    sig_ = cc.get("analysed_random_significance_2026_10", {})
    rows.append({"part": "ClinVar check (labels only, no scores)",
                 "quantity": "analysed random-labelled variants present in October 2026, ClinVar classification: " +
                             "; ".join(f"{('no classification' if k == '-' else k)} {v:,}" for k, v in sig_.items())})
    rows.append({"part": "ClinVar check (labels only, no scores)",
                 "quantity": f"totals: random-labelled {cc['compass_main_snvs']['random_snv']:,} "
                             f"({cc['analysed_snvs']['random_snv']:,} analysed); ExAC subset "
                             f"{cc['compass_main_snvs']['exac']:,} ({cc['analysed_snvs']['exac']:,} analysed)"})
    out.append(("COMPASS without its ExAC subset (additional analysis) and the ClinVar check of its random variants",
                "Added after a review comparing the manuscript with the analysis code; defined in the repository "
                "before it was computed. The ExAC subset was chosen by the COMPASS authors by presence in a population "
                "database. Its own estimate is descriptive and is not pooled on its own; its variants are part of "
                "COMPASS in the planned pooled estimate (Supplementary Table 12) and are left out only in the additional "
                "pooled estimate here. The random variants are the authors' random label minus "
                "alleles present in the ClinVar release of 2025-06-15; the ClinVar check, on labels alone, shows that "
                "none of them was in the ClinVar summaries of July 2024 or July 2025, and that those analysed and "
                "present in October 2026 carry no P/LP or B/LB classification. Counts beside each AUROC are the "
                "variants that tool scored. " + DEF,
                "results/part2/compass_subsets_*.tsv/json; results/part2/compass_clinvar_check.json", pd.DataFrame(rows)))
    return out


def minus_signs(df: pd.DataFrame) -> pd.DataFrame:
    """Comparison labels with a true minus sign (U+2212): 'AVI - SpliceAI' -> 'AVI − SpliceAI'."""
    out = df.copy()
    for col in out.columns:
        if out[col].dtype == object:
            out[col] = out[col].map(lambda v: v.replace("AVI - ", "AVI − ").replace("mutant - WT", "mutant − WT")
                                    if isinstance(v, str) else v)
    return out


NO_FIXED = ("likelihood", "threshold", "target", "score", "tau", "prior", "I²")


def fixed_format(ws, hdr: int, ratio_sheet: bool) -> None:
    """Four-decimal display for metric cells (values are already rounded)."""
    heads = {c.column: str(c.value or "") for c in ws[hdr]}
    for row in ws.iter_rows(min_row=hdr + 1):
        for c in row:
            h = heads.get(c.column, "")
            if not isinstance(c.value, float):
                continue
            if any(k in h for k in NO_FIXED):
                continue
            if ratio_sheet and not h.startswith("share"):
                continue
            c.number_format = "0.0000"


def main() -> int:
    head = subprocess.run(["git", "rev-parse", "--short", "HEAD"], capture_output=True, text=True).stdout.strip()
    topic = tables()
    if sorted(NUMBERING_ORDER) != list(range(1, len(topic) + 1)):
        sys.exit("NUMBERING_ORDER does not cover every table once")
    tabs = [topic[i - 1] for i in NUMBERING_ORDER]
    contents = pd.DataFrame([{"table": f"Supplementary Table {i}", "title": t, "legend": leg, "source": src,
                              "rows": len(df)} for i, (t, leg, src, df) in enumerate(tabs, 1)])
    definitions = pd.DataFrame(DEFINITIONS, columns=["term", "definition"])
    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        intro = pd.DataFrame({"": [
            f"Supplementary tables for: {TITLE}.",
            f"Built from the aggregate result tables of the analysis repository at commit {head}. No per-variant tool "
            "score is included. Definitions of sets, regions, tools, abbreviations and notation follow the list of "
            "tables."]})
        intro.to_excel(xw, sheet_name="Contents", index=False, header=False, startrow=0)
        contents.to_excel(xw, sheet_name="Contents", index=False, startrow=3)
        definitions.to_excel(xw, sheet_name="Contents", index=False, startrow=3 + len(contents) + 3)
        for i, (t, leg, src, df) in enumerate(tabs, 1):
            name = f"S{i}"
            pd.DataFrame({"": [f"Supplementary Table {i}. {t}", leg, f"Source: {src}"]}).to_excel(
                xw, sheet_name=name, index=False, header=False)
            minus_signs(df.drop(columns=[c for c in df.columns if c.startswith("_")])).to_excel(
                xw, sheet_name=name, index=False, startrow=4)
    wb = load_workbook(OUT)
    def_row = 3 + len(contents) + 3 + 1
    for ws in wb.worksheets:
        ws["A1"].font = Font(bold=True)
        hdr = 5 if ws.title != "Contents" else 4
        for cell in ws[hdr]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if ws.title == "Contents":
            for cell in ws[def_row]:
                cell.font = Font(bold=True)
            for row in ws.iter_rows(min_row=def_row + 1):
                if row[0].value and not row[1].value:
                    row[0].font = Font(bold=True)
        else:
            title = str(ws["A1"].value)
            fixed_format(ws, hdr, ratio_sheet=("Calibration" in title or "Published thresholds" in title))
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col[hdr - 1:] if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
        ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
    wb.save(OUT)
    for i, (t, _, _, df) in enumerate(tabs, 1):
        print(f"S{i}\t{len(df)} rows\t{t}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
